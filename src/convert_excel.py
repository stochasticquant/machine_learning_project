"""
LAEI 2019 -- one-off Excel -> CSV conversion.

The source workbooks are far too large to open with pandas.read_excel
(the road-link workbook alone is 345 MB and expands to many GB in memory).
This module streams each sheet once with openpyxl in read-only mode and
writes a flat CSV. Run it ONCE; everything downstream reads the CSVs.

Three source-specific hazards are handled here, all verified against the
2019 files:

  1. The workbooks declare a bogus 1x1 sheet dimension. openpyxl's
     read-only mode believes it and yields a single cell per row, so
     ws.reset_dimensions() must be called on every sheet.

  2. The PM10 / PM2.5 link sheets carry an extra `pm-source` column and
     hold THREE rows per TOID (exhaust, brake wear, tyre wear), where the
     NOx and CO2 sheets hold one. Emissions must be SUMMED across sources;
     keying a dict on TOID and assigning would silently discard ~2/3 of
     the PM mass.

  3. Column headers carry leading/trailing whitespace (' AADT Taxi ').

Usage:
    python src/convert_excel.py
"""

from __future__ import annotations

import csv
import sys
import time
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
INTERIM = ROOT / "data" / "interim"
PROCESSED = ROOT / "data" / "processed"

GRID_XLSX = INTERIM / "LAEI-2019-Emissions-Summary-including-Forecast.xlsx"
FLOWS_XLSX = INTERIM / "laei-2019-major-roads-vkm-flows-speeds.xlsx"
LINKS_XLSX = INTERIM / "LAEI2019-nox-pm-co2-major-roads-link-emissions.xlsx"

# Expected row counts, verified by inspection of the 2019 release.
EXPECT = {
    "grid_all_years.csv": 699_120,
    "link_features.csv": 79_437,
    "link_targets.csv": 79_439,
}


def _sheet(wb, name):
    """Return a read-only worksheet with its bogus dimensions reset (hazard 1)."""
    ws = wb[name]
    ws.reset_dimensions()
    return ws


def _header(row_iter):
    """Consume and clean the header row (hazard 3)."""
    return [str(c).strip() if c is not None else "" for c in next(row_iter)]


def convert_grid(out_csv: Path) -> int:
    """Emissions by Grid ID -> CSV. All 5 years, all 19 pollutants, 1 km cells."""
    wb = openpyxl.load_workbook(GRID_XLSX, read_only=True)
    ws = _sheet(wb, "Emissions by Grid ID")
    rows = ws.iter_rows(values_only=True)
    header = _header(rows)
    n = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for row in rows:
            w.writerow(row)
            n += 1
    wb.close()
    return n


def convert_link_features(out_csv: Path) -> int:
    """Major-road AADT / speed / VKM per link -> CSV. This is the feature table."""
    wb = openpyxl.load_workbook(FLOWS_XLSX, read_only=True)
    ws = _sheet(wb, "laei-2019-major-roads")
    rows = ws.iter_rows(values_only=True)
    header = _header(rows)
    n = 0
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        for row in rows:
            w.writerow(row)
            n += 1
    wb.close()
    return n


def convert_link_targets(out_csv: Path) -> int:
    """
    Per-link 2019 emissions -> one row per TOID with four target columns.

    PM10 and PM2.5 are accumulated across `pm-source` (hazard 2); the running
    total uses `+=` deliberately. `pm_sources` is returned in the log so the
    3-way split can be confirmed rather than assumed.
    """
    sheets = [
        ("NOx Road Link Emissions", "nox"),
        ("PM10 Road Link Emissions", "pm10"),
        ("PM2.5 Road Link Emissions", "pm25"),
        ("CO2 Road Link Emissions", "co2"),
    ]
    wb = openpyxl.load_workbook(LINKS_XLSX, read_only=True)
    agg: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    pm_sources: set[tuple[str, str]] = set()

    for sheet_name, tag in sheets:
        t0 = time.time()
        ws = _sheet(wb, sheet_name)
        rows = ws.iter_rows(values_only=True)
        header = _header(rows)
        i_toid = header.index("TOID")
        i_total = header.index("Road-Total-2019")
        i_src = header.index("pm-source") if "pm-source" in header else None

        n = 0
        for row in rows:
            agg[row[i_toid]][tag] += (row[i_total] or 0)
            if i_src is not None:
                pm_sources.add((tag, row[i_src]))
            n += 1
        print(f"    {sheet_name:<32} {n:>7,} rows  {time.time() - t0:5.0f}s", flush=True)

    wb.close()
    print(f"    pm-source values: {sorted(pm_sources)}", flush=True)

    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["TOID", "nox", "pm10", "pm25", "co2"])
        for toid, vals in agg.items():
            w.writerow([toid, vals["nox"], vals["pm10"], vals["pm25"], vals["co2"]])
    return len(agg)


def main() -> int:
    PROCESSED.mkdir(parents=True, exist_ok=True)
    jobs = [
        ("link_features.csv", convert_link_features),
        ("grid_all_years.csv", convert_grid),
        ("link_targets.csv", convert_link_targets),
    ]
    failures = []
    for name, fn in jobs:
        out = PROCESSED / name
        print(f"[{name}] converting ...", flush=True)
        t0 = time.time()
        n = fn(out)
        size_mb = out.stat().st_size / 1e6
        expected = EXPECT[name]
        ok = "OK" if n == expected else f"MISMATCH (expected {expected:,})"
        if n != expected:
            failures.append(name)
        print(
            f"[{name}] {n:,} rows  {size_mb:,.0f} MB  "
            f"{time.time() - t0:,.0f}s  {ok}\n",
            flush=True,
        )

    if failures:
        print("FAILED row-count checks:", ", ".join(failures))
        return 1
    print("All conversions complete and row counts verified.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
